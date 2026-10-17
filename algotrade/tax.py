"""Realised P&L for a financial year (April to March), split the way Indian income tax treats it, from the orders the
dashboard saved (orderbook.order_history). An estimate to help with filing, not a tax statement: Kotak's own
tax P&L report (Kotak Neo > Reports) is the official one.

- Each sell is matched against earlier buys of the same instrument, oldest first (FIFO).
- Equity: bought and sold the same day (or MIS) = intraday (speculative); held up to 12 months = short-term
  capital gain (STCG); longer = long-term (LTCG).
- Options (F&O) = non-speculative business income.
- A sell with no saved buy before it (bought before the app kept history) is listed with its cost unknown.
- Brokerage, STT and other charges are not included."""
import io
import re
from datetime import date, datetime

from . import orderbook

OPTION = re.compile(r"\d(CE|PE)$")
CATEGORIES = ("STCG", "LTCG", "Intraday", "F&O")


def fy_range(fy: int) -> tuple[date, date]:
    return date(fy, 4, 1), date(fy + 1, 3, 31)


def current_fy(today: date | None = None) -> int:
    today = today or date.today()
    return today.year if today.month >= 4 else today.year - 1


def _fills() -> list[dict]:
    rows = []
    for o in orderbook.history(orderbook.HISTORY_DAYS):
        qty, px = o["filled_qty"] or 0, o["avg_price"]
        if qty <= 0 or not px:
            continue
        rows.append({"symbol": o["symbol"].removesuffix("-EQ"), "side": o["side"], "qty": float(qty), "price": float(px),
                     "time": datetime.fromisoformat(o["time"]), "product": o["product"] or "",
                     "option": bool(OPTION.search(o["symbol"] or ""))})
    return sorted(rows, key=lambda r: r["time"])


def report(fy: int | None = None) -> dict:
    fy = current_fy() if fy is None else fy
    start, end = fy_range(fy)
    lots: dict[str, list[dict]] = {}   # symbol -> open buy lots, oldest first
    realised, unmatched = [], []
    fills = _fills()
    for f in fills:
        if f["side"] == "BUY":
            lots.setdefault(f["symbol"], []).append({**f})
            continue
        left = f["qty"]
        while left > 0 and lots.get(f["symbol"]):
            lot = lots[f["symbol"]][0]
            q = min(left, lot["qty"])
            days = (f["time"].date() - lot["time"].date()).days
            if f["option"]:
                cat = "F&O"
            elif days == 0 or "MIS" in (f["product"], lot["product"]):
                cat = "Intraday"
            else:
                cat = "LTCG" if days > 365 else "STCG"
            if start <= f["time"].date() <= end:
                realised.append({"symbol": f["symbol"], "category": cat, "qty": q, "buy_date": lot["time"].date().isoformat(),
                                 "buy_price": lot["price"], "sell_date": f["time"].date().isoformat(), "sell_price": f["price"],
                                 "buy_value": round(q * lot["price"], 2), "sell_value": round(q * f["price"], 2),
                                 "pnl": round(q * (f["price"] - lot["price"]), 2), "days": days})
            lot["qty"] -= q
            left -= q
            if lot["qty"] <= 1e-9:
                lots[f["symbol"]].pop(0)
        if left > 1e-9 and start <= f["time"].date() <= end:
            unmatched.append({"symbol": f["symbol"], "qty": left, "sell_date": f["time"].date().isoformat(),
                              "sell_price": f["price"], "sell_value": round(left * f["price"], 2),
                              "category": "F&O" if f["option"] else "Equity (cost unknown)"})
    totals = {c: {"pnl": round(sum(r["pnl"] for r in realised if r["category"] == c), 2),
                  "trades": sum(1 for r in realised if r["category"] == c)} for c in CATEGORIES}
    first = fills[0]["time"].date().isoformat() if fills else None
    return {"fy": fy, "label": f"FY {fy}-{str(fy + 1)[-2:]}", "from": start.isoformat(), "to": end.isoformat(),
            "realised": realised, "unmatched": unmatched, "totals": totals,
            "total_pnl": round(sum(r["pnl"] for r in realised), 2), "history_from": first,
            "open_lots": [{"symbol": s, "qty": sum(l["qty"] for l in ls), "since": ls[0]["time"].date().isoformat()}
                          for s, ls in lots.items() if ls]}


def xlsx(fy: int | None = None) -> bytes:
    """The report as an Excel workbook: Summary, Realised and Cost unknown sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = report(fy)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"AlgoTrade realised P&L · {r['label']} ({r['from']} to {r['to']})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Estimate from the orders the dashboard saved"
               + (f" (since {r['history_from']})" if r["history_from"] else "") + "; charges not included. "
               "Kotak's tax P&L report is the official one."])
    ws.append([])
    ws.append(["Category", "Trades", "P&L (₹)"])
    for c in CATEGORIES:
        ws.append([c, r["totals"][c]["trades"], r["totals"][c]["pnl"]])
    ws.append(["Total", sum(t["trades"] for t in r["totals"].values()), r["total_pnl"]])
    for row in (4, 9):
        for cell in ws[row]:
            cell.font = Font(bold=True)
    cols = ["symbol", "category", "qty", "buy_date", "buy_price", "sell_date", "sell_price", "buy_value", "sell_value",
            "pnl", "days"]
    sh = wb.create_sheet("Realised")
    sh.append([c.replace("_", " ").capitalize() for c in cols])
    for x in r["realised"]:
        sh.append([x[c] for c in cols])
    sh2 = wb.create_sheet("Cost unknown")
    ucols = ["symbol", "category", "qty", "sell_date", "sell_price", "sell_value"]
    sh2.append([c.replace("_", " ").capitalize() for c in ucols])
    for x in r["unmatched"]:
        sh2.append([x[c] for c in ucols])
    for s in (sh, sh2):
        for cell in s[1]:
            cell.font = Font(bold=True)
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
