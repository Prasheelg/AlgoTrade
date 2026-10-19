"""Second feature batch: settings, portfolio risk, events filter, daily summary, scorecard + dashboard backtest,
Telegram commands, phone push, market overview, tax report and the trade journal. Kotak, Yahoo, NSE and Telegram
are faked."""
import io
import json
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest

import main
from algotrade import (backtest_job, broker, config, events, journal, kotak, market, markets, orderbook, orders,
                       performance, push, risk, scanner, settings, state, summary, tax, telegram_bot, tgcommands,
                       universe, web)
from conftest import Recorder, make_signal

IST = ZoneInfo("Asia/Kolkata")


@pytest.fixture
def msgs(monkeypatch):
    rec = Recorder(True)
    monkeypatch.setattr(telegram_bot, "send_message", rec)
    return rec


@pytest.fixture
def cfg(monkeypatch):
    """Let settings.apply() change config / markets; monkeypatch puts everything back afterwards."""
    for name in settings.EDITABLE:
        monkeypatch.setattr(config, name, getattr(config, name))
    monkeypatch.setattr(markets, "MARKETS", dict(markets.MARKETS))
    monkeypatch.setattr(settings, "_DEFAULTS", {**settings._DEFAULTS, "CAPITAL": 100000.0, "MAX_OPEN_RISK_PCT": 6.0})


# ---- settings -----------------------------------------------------------------------------------

def test_settings_override_env_until_reset(cfg, msgs):
    out = settings.update({"CAPITAL": 200000, "MAX_OPEN_RISK_PCT": 4, "DAILY_SUMMARY": False}, by="test")
    assert config.CAPITAL == 200000 and markets.get("IN").capital == 200000 and config.DAILY_SUMMARY is False
    assert {x["name"]: x["changed"] for x in out}["CAPITAL"] is True
    assert "Settings changed" in msgs.calls[0][0][0]
    settings.update({"CAPITAL": None}, by="test")
    assert config.CAPITAL == 100000.0 and markets.get("IN").capital == 100000.0


def test_other_process_picks_up_changes(cfg, msgs):
    state.set_setting(settings.KEY, {"RISK_PER_TRADE_PCT": 0.5})
    settings.apply()
    assert config.RISK_PER_TRADE_PCT == 0.5


@pytest.mark.parametrize("changes,msg", [({"CAPITAL": 10}, "between"), ({"OPTIONS_MAX_LOTS": 1.5}, "whole number"),
                                         ({"DAILY_SUMMARY": "yes"}, "on or off"), ({"KOTAK_MPIN": "1"}, "Can't change"),
                                         ({"CAPITAL": "abc"}, "a number")])
def test_bad_settings_change_nothing(cfg, changes, msg):
    with pytest.raises(ValueError, match=msg):
        settings.update(changes, by="test")
    assert state.get_setting(settings.KEY, None)[0] is None


def test_settings_api_needs_trading_access(cfg, msgs):
    c = web.app.test_client()
    r = c.post("/api/settings", json={"changes": {"CAPITAL": 150000}}, environ_base={"REMOTE_ADDR": "192.168.1.9"})
    assert r.status_code == 401 and config.CAPITAL == 100000.0
    r = c.post("/api/settings", json={"changes": {"CAPITAL": 150000}})
    assert r.status_code == 200 and config.CAPITAL == 150000
    assert c.get("/api/settings").get_json()["can_change"] is True


# ---- portfolio risk ------------------------------------------------------------------------------

def test_open_risk_caps_auto_trade(cfg, monkeypatch, msgs):
    monkeypatch.setattr(config, "MAX_OPEN_RISK_PCT", 0.5)          # ₹500 on ₹1,00,000
    monkeypatch.setattr(config, "MAX_ORDER_VALUE", 1e6)
    placed = []
    monkeypatch.setattr(broker, "call", lambda m, **kw: {"data": []} if m == "order_report" else placed.append(kw) or {"nOrdNo": "1"})
    monkeypatch.setattr(orders, "_market_open", lambda: True)
    state.set_setting(broker.TRADING_KEY, {"enabled": True})
    s = make_signal(symbol="RELIANCE")   # limit 99, stop 96: ₹1,000 risk for 333 shares
    state.record([s])
    orders.enter([s.key])
    t = orders._get(s.key)
    assert t["state"] == "SKIPPED" and "portfolio risk limit" in t["note"] and placed == []


def test_risk_report(monkeypatch):
    orders._save("k1", symbol="TCS", state=orders.OPEN, qty=10, entry=100.0, fill_price=101.0, stop=95.0, created_day="d")
    orders._save("k2", symbol="INFY", state=orders.CLOSED, qty=10, entry=100.0, stop=95.0, created_day="d")
    monkeypatch.setattr(kotak, "quotes", lambda syms: {"TCS": {"ltp": 110.0}})
    monkeypatch.setattr(universe, "sectors", lambda: {"TCS": "IT", "HDFCBANK": "Financial Services"})
    r = risk.report(holdings=[{"symbol": "HDFCBANK", "qty": 5, "avg": 1500.0, "ltp": 1600.0, "value": 8000.0},
                              {"symbol": "TCS", "qty": 10, "avg": 101.0, "ltp": 110.0, "value": 1100.0}])
    assert r["at_risk"] == 60.0 and r["at_risk_now"] == 150.0 and r["unprotected"] == 8000.0
    assert [p["symbol"] for p in r["positions"]] == ["TCS", "HDFCBANK"]       # TCS counted once (app trade)
    assert r["largest_sector"]["sector"] == "Financial Services" and r["limit"] == 6000.0


# ---- events filter ---------------------------------------------------------------------------------

def test_events_filter_drops_signals_before_results(monkeypatch):
    soon = date.today() + timedelta(days=2)
    monkeypatch.setattr(events, "_cache", {})

    class T:
        def __init__(self, sym):
            self.calendar = {"Earnings Date": [soon]} if sym.startswith("A") else {}
    monkeypatch.setattr(events.yf, "Ticker", T)
    assert "results in 2 days" in events.soon("AAA") and events.soon("BBB") is None
    monkeypatch.setattr(scanner, "fetch_ohlcv", lambda *a, **k: {"AAA": 1, "BBB": 2})
    monkeypatch.setattr(scanner, "fetch_index", lambda *a, **k: None)
    monkeypatch.setattr(scanner, "fill_missing_days", lambda d, i, *a: (d, i))
    monkeypatch.setattr(scanner, "completed_bars", lambda df, *a: df)
    monkeypatch.setattr(scanner, "add_all", lambda df, idx: df)
    monkeypatch.setattr(scanner, "evaluate", lambda sym, df, **kw: make_signal(symbol=sym))
    assert [s.symbol for s in scanner.scan(["AAA", "BBB"], filters=["events"])] == ["BBB"]
    assert len(scanner.scan(["AAA", "BBB"])) == 2


# ---- daily summary ----------------------------------------------------------------------------------

def test_summary_once_after_close(msgs, monkeypatch):
    monkeypatch.setattr(summary, "_account", lambda: ["Cash available ₹5,000"])
    s = make_signal(symbol="TCS")
    state.record([s])
    state.update_signal(s.key, status="ACTIVE", ltp=100.0, fill_price=99.0)
    fri = lambda h, m=0: datetime(2026, 10, 16, h, m, tzinfo=IST)
    assert summary.daily(fri(15, 0)) is False
    assert summary.daily(fri(15, 50)) is True and summary.daily(fri(16, 30)) is False
    text = msgs.calls[0][0][0]
    assert "Daily summary" in text and "Cash available ₹5,000" in text and "TCS +1.0%" in text
    assert summary.daily(datetime(2026, 10, 17, 16, 0, tzinfo=IST)) is False    # Saturday


def test_summary_off_by_setting(msgs, monkeypatch):
    monkeypatch.setattr(config, "DAILY_SUMMARY", False)
    assert summary.daily(datetime(2026, 10, 16, 16, 0, tzinfo=IST)) is False and msgs.calls == []


# ---- scorecard + dashboard backtest -----------------------------------------------------------------

def test_backtest_summary_and_scorecard(monkeypatch):
    trades = [{"symbol": "A", "side": "BUY", "exit_time": f"2026-01-{d:02d}", "R": r + 0.1, "net_R": r,
               "strategies": "EMA 9/21 Cross" + (", RSI Pullback" if d % 2 else "")} for d, r in
              [(k, 1.0 if k % 3 == 0 else -0.5) for k in range(1, 31)]]
    s = backtest_job.summarize(trades)
    assert s["net"]["trades"] == 30 and s["net"]["total_r"] == 0.0 and s["equity"][-1]["net"] == 0.0
    by = {r["group"]: r for r in s["by_strategy"]}
    assert by["EMA 9/21 Cross"]["trades"] == 30 and by["RSI Pullback"]["trades"] == 15
    backtest_job._path("IN").write_text(json.dumps({"finished": "x", "params": {}, "universe": 1, **s}))
    sc = {r["name"]: r for r in performance.scorecard("IN")["rows"]}
    assert (sc["EMA 9/21 Cross"]["verdict"], sc["EMA 9/21 Cross"]["basis"]) == ("Watch", "backtest")
    assert sc["RSI Pullback"]["verdict"] == "Too few trades"


def test_switching_a_strategy_off_reaches_the_scanner(msgs, monkeypatch):
    c = web.app.test_client()
    r = c.post("/api/scorecard", json={"market": "IN", "name": "EMA 9/21 Cross", "off": True})
    assert r.status_code == 200 and r.get_json()["off"] == ["EMA 9/21 Cross"]
    kw = main.scanner_kw("IN")
    assert "EMA 9/21 Cross" not in kw["strategies"] and "MACD Trend Cross" in kw["strategies"]
    assert c.post("/api/scorecard", json={"market": "IN", "name": "Nope", "off": True}).status_code == 400
    r = c.post("/api/scorecard", json={"market": "IN", "name": "MACD Trend Cross", "off": True},
               environ_base={"REMOTE_ADDR": "192.168.1.9"})
    assert r.status_code == 401


def test_backtest_job_validates(monkeypatch):
    with pytest.raises(ValueError, match="Period"):
        backtest_job.start("IN", ["EMA 9/21 Cross"], 1, [], "3w")
    with pytest.raises(ValueError, match="strategy"):
        backtest_job.start("IN", [], 1, [], "5y")
    with pytest.raises(ValueError, match="between 1 and 1"):
        backtest_job.start("IN", ["EMA 9/21 Cross"], 2, [], "5y")


def test_backtest_job_runs_in_background(monkeypatch):
    import pandas as pd
    from algotrade import backtest
    monkeypatch.setattr("algotrade.universe.get_universe", lambda **kw: ["A", "B"])
    got = {}

    def run(syms, tf, **kw):
        got.update(kw)
        return pd.DataFrame([{"symbol": "A", "side": "BUY", "exit_time": "2026-01-02", "R": 1.0, "net_R": 0.9,
                              "strategies": "EMA 9/21 Cross"}])
    monkeypatch.setattr(backtest, "run", run)
    backtest_job.start("IN", ["EMA 9/21 Cross"], 1, ["trend", "events"], "2y")
    for _ in range(100):
        if not backtest_job._job["running"]:
            break
        import time as _t
        _t.sleep(0.02)
    last = backtest_job.status("IN")["last"]
    assert last["net"]["trades"] == 1 and got["filters"] == ["trend"] and got["tag_strategies"] is True


# ---- Telegram commands --------------------------------------------------------------------------------

def test_commands(monkeypatch, msgs):
    monkeypatch.setattr(kotak, "tokens", lambda: {"TCS": "1"})
    state.set_setting(broker.TRADING_KEY, {"enabled": True})
    monkeypatch.setattr(broker, "login", lambda force=False: None)
    assert "AlgoTrade commands" in tgcommands.answer("/help@MyBot")
    assert "switched <b>OFF</b>" in tgcommands.answer("/stop") and not broker.trading_enabled()
    assert "Added TCS" in tgcommands.answer("/watch tcs")
    assert "Alert set: TCS above ₹2,500.00" in tgcommands.answer("/alert TCS above ₹2,500")
    assert "TCS above" in tgcommands.answer("/alerts")
    assert "⚠️" in tgcommands.answer("/alert TCS near 1")
    assert tgcommands.answer("/unknown") is None and tgcommands.answer("hi") is None
    assert "on" not in tgcommands.COMMANDS and "start_trading" not in tgcommands.COMMANDS


def test_poll_answers_only_our_group(monkeypatch, msgs):
    monkeypatch.setattr(config, "TELEGRAM_BOT_TOKEN", "t")
    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "-100")
    monkeypatch.setattr(config, "TELEGRAM_COMMANDS", True)
    seen = []

    class R:
        ok = True

        def __init__(self, ups):
            self.ups = ups

        def json(self):
            return {"result": self.ups}
    ups = [{"update_id": 5, "message": {"chat": {"id": -100}, "text": "/status"}},
           {"update_id": 6, "message": {"chat": {"id": 999}, "text": "/stop"}}]
    monkeypatch.setattr(tgcommands.requests, "get", lambda url, params, timeout: seen.append(params) or R(ups))
    assert tgcommands.poll_once(0) == 1 and "Auto-trade" in msgs.calls[0][0][0]
    tgcommands.poll_once(0)
    assert seen[-1]["offset"] == 7


# ---- push --------------------------------------------------------------------------------------------

def test_push_subscribe_send_and_cleanup(monkeypatch):
    sub = {"endpoint": "https://push.example/abc", "keys": {"p256dh": "x", "auth": "y"}}
    c = web.app.test_client()
    assert len(c.get("/api/push/key").get_json()["key"]) > 60
    assert c.post("/api/push/subscribe", json={"subscription": sub}).get_json()["devices"] == 1
    assert c.post("/api/push/subscribe", json={"subscription": {"endpoint": "http://x"}}).status_code == 400
    sent = []
    import pywebpush

    class Gone(pywebpush.WebPushException):
        pass

    def webpush(s, data, **kw):
        sent.append(json.loads(data))
        if len(sent) > 1:
            raise Gone("gone", response=type("R", (), {"status_code": 410})())
    monkeypatch.setattr(pywebpush, "webpush", webpush)
    assert push.send("Hi", "there", "/x") == 1 and sent[0] == {"title": "Hi", "body": "there", "url": "/x"}
    assert push.send("Hi", "again") == 0 and push.devices() == []       # 410: subscription removed


# ---- market overview ----------------------------------------------------------------------------------

def test_market_overview(monkeypatch):
    monkeypatch.setattr(market, "_cache", {"at": 0.0, "data": None, "flows_at": 0.0, "flows": {"date": "x"},
                                           "universe": ["A", "B", "C"], "universe_day": date.today()})
    monkeypatch.setattr(market, "FLOWS_SEC", 10 ** 9)
    monkeypatch.setattr(market.time, "monotonic", lambda: 100.0)
    monkeypatch.setattr(kotak, "index_quotes", lambda names: {"Nifty 50": {"ltp": 101.0, "prev_close": 100.0},
                                                              "Nifty IT": {"ltp": 98.0, "prev_close": 100.0}})
    monkeypatch.setattr(kotak, "quotes", lambda syms: {"A": {"ltp": 110.0, "prev_close": 100.0},
                                                       "B": {"ltp": 95.0, "prev_close": 100.0},
                                                       "C": {"ltp": 100.0, "prev_close": None}})
    monkeypatch.setattr(universe, "sectors", lambda: {"A": "IT", "B": "IT"})
    o = market.overview()
    assert o["indices"] == [{"name": "Nifty 50", "ltp": 101.0, "change": 1.0, "change_pct": 1.0}]
    assert o["sectors"][0]["name"] == "IT" and [g["symbol"] for g in o["gainers"]] == ["A", "B"]
    assert o["losers"][0]["symbol"] == "B" and o["breadth"] == {"up": 1, "down": 1, "total": 2}
    assert o["industries"] == [{"name": "IT", "stocks": 2, "change_pct": 2.5}] and o["flows"] == {"date": "x"}


# ---- tax report --------------------------------------------------------------------------------------

def fill(oid, sym, side, qty, px, when, product="CNC"):
    return {"order_id": oid, "symbol": sym, "side": side, "qty": qty, "filled_qty": qty, "avg_price": px,
            "state": "COMPLETED", "product": product, "time": when}


def test_tax_fifo_categories(monkeypatch):
    t0 = date.today()
    d = lambda days: (t0 - timedelta(days=days)).strftime("%d-%b-%Y") + " 10:00:00"
    orderbook.save_history([
        fill("1", "TCS-EQ", "BUY", 10, 100.0, d(500)), fill("2", "TCS-EQ", "BUY", 10, 120.0, d(100)),
        fill("3", "TCS-EQ", "SELL", 15, 150.0, d(1)),                                   # 10 LTCG + 5 STCG
        fill("4", "INFY-EQ", "BUY", 5, 1000.0, d(2).replace("10:00", "09:30"), "MIS"),
        fill("5", "INFY-EQ", "SELL", 5, 990.0, d(2).replace("10:00", "15:00"), "MIS"),   # intraday
        fill("6", "NIFTY26O1923000CE", "BUY", 65, 50.0, d(3)), fill("7", "NIFTY26O1923000CE", "SELL", 65, 70.0, d(2)),
        fill("8", "WIPRO-EQ", "SELL", 3, 500.0, d(1)),                                  # bought before history
    ])
    r = tax.report(tax.current_fy(t0 - timedelta(days=1)))
    tot = {k: v["pnl"] for k, v in r["totals"].items()}
    in_fy = lambda days: tax.fy_range(r["fy"])[0] <= t0 - timedelta(days=days)
    assert tot["LTCG"] == 500.0 and tot["STCG"] == 150.0
    assert tot["Intraday"] == (-50.0 if in_fy(2) else 0) and tot["F&O"] == (1300.0 if in_fy(2) else 0)
    assert r["unmatched"][0]["symbol"] == "WIPRO" and r["open_lots"] == [{"symbol": "TCS", "qty": 5.0,
                                                                         "since": (t0 - timedelta(days=100)).isoformat()}]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(tax.xlsx(r["fy"])))
    assert wb.sheetnames == ["Summary", "Realised", "Cost unknown"] and wb["Realised"].max_row == len(r["realised"]) + 1
    resp = web.app.test_client().get(f"/api/tax.xlsx?fy={r['fy']}")
    assert resp.status_code == 200 and "attachment" in resp.headers["Content-Disposition"]


def test_fy_boundaries():
    assert tax.current_fy(date(2026, 3, 31)) == 2025 and tax.current_fy(date(2026, 4, 1)) == 2026
    assert tax.fy_range(2026) == (date(2026, 4, 1), date(2027, 3, 31))


# ---- journal -------------------------------------------------------------------------------------------

def test_journal_and_tag_breakdown():
    s = make_signal(symbol="TCS")
    state.record([s])
    c = web.app.test_client()
    r = c.post("/api/journal", json={"key": s.key, "note": "late entry", "tags": ["Entered late", "Entered late", " "]})
    assert r.status_code == 200 and r.get_json()["tags"] == ["Entered late"]
    assert c.post("/api/journal", json={"key": "nope", "note": "x"}).status_code == 400
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 100
    r = c.post("/api/journal/image", data={"key": s.key, "image": (io.BytesIO(png), "c.png", "image/png")},
               content_type="multipart/form-data")
    img = r.get_json()["image"]
    assert r.status_code == 200 and img.endswith(".png") and c.get(f"/journal-img/{img}").data == png
    r = c.post("/api/journal/image", data={"key": s.key, "image": (io.BytesIO(b"x"), "a.gif", "image/gif")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    state.update_signal(s.key, status="SL_HIT", exit_price=96.0, closed_at=date.today().isoformat() + "T10:00:00",
                        fill_price=99.0)
    bd = {b["name"]: b for b in performance.report("IN")["breakdowns"]}
    assert [r["group"] for r in bd["Tag"]["rows"]] == ["Entered late"]
    assert c.post("/api/journal/image/delete", json={"key": s.key}).get_json()["image"] is None
